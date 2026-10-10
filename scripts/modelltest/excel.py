#!/usr/bin/env python3
"""Bygger sammanställningen: resultat + maskinella mått + domarens betyg -> xlsx.

  python3 scripts/modelltest/excel.py [utfil.xlsx]

Läser  resultat/<körning>/<kod>.json   (glosorna, blindade)
       facit/facit.json                (kod -> modell, latens, kostnad)
       domar/<körning>.json            (domarens betyg, se domare-instruktion.md)

Facit läses FÖRST här, i sammanställningen – aldrig under bedömningen.
"""
import json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from matt import matt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HAR = os.path.dirname(os.path.abspath(__file__))
UT, FACIT, DOMAR = (os.path.join(HAR, d) for d in ("resultat", "facit", "domar"))
VIKTER = {"sakfel": 0.40, "relevans": 0.30, "prio": 0.20, "bojning": 0.10}

RUBRIK = Font(bold=True, color="FFFFFF")
FYLL = PatternFill("solid", fgColor="2C3C63")
FET = Font(bold=True)

def las(p, standard=None):
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else standard

def viktat(betyg):
    if not betyg: return None
    return round(sum(VIKTER[k] * betyg.get(k, 0) for k in VIKTER), 2)

def kolbredd(ws, bredder):
    for i, b in enumerate(bredder, 1): ws.column_dimensions[get_column_letter(i)].width = b

def main():
    utfil = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HAR, "modelltest.xlsx")
    facit = las(os.path.join(FACIT, "facit.json"), {}) or {}
    wb = Workbook(); wb.remove(wb.active)
    sammanst = {}   # modell -> lista av dict per körning

    for kid in sorted(os.listdir(UT)):
        mapp = os.path.join(UT, kid)
        if not os.path.isdir(mapp): continue
        domar = (las(os.path.join(DOMAR, kid + ".json"), {}) or {}).get("domar", {})
        rader = []
        for f in sorted(os.listdir(mapp)):
            kod = f[:-5]
            d = las(os.path.join(mapp, f))
            fk = facit.get(f"{kid}/{kod}", {})
            dom = domar.get(kod, {})
            rader.append({"kod": kod, "data": d, "matt": matt(d), "facit": fk, "dom": dom,
                          "modell": fk.get("modell", "(okänd)"), "poang": viktat(dom.get("betyg"))})
        rader.sort(key=lambda r: (-(r["poang"] or 0), r["modell"]))

        ws = wb.create_sheet(kid[:31])
        meta = rader[0]["data"]["meta"] if rader else {}
        ws["A1"] = f'{meta.get("tema","")} · {meta.get("sprak","")} · {meta.get("antal","")} glosor'
        ws["A1"].font = Font(bold=True, size=13)

        # --- översta blocket: en rad per modell ---
        kol = ["modell", "poäng", "sakfel", "relevans", "prio", "böjning", "levererat", "täckning",
               "dubbletter", "tom böjning %", "böjn. ekar", "prio 1/2/3", "prio1 svans",
               "sekunder", "kostnad $", "kommentar"]
        for j, k in enumerate(kol, 1):
            c = ws.cell(row=3, column=j, value=k); c.font = RUBRIK; c.fill = FYLL
        for i, r in enumerate(rader, 4):
            m, b, dm = r["matt"], r["dom"].get("betyg", {}), r["facit"]
            ws.cell(row=i, column=1, value=r["modell"]).font = FET
            for j, v in enumerate([
                r["poang"], b.get("sakfel"), b.get("relevans"), b.get("prio"), b.get("bojning"),
                m["levererat"], m["täckning"], m["dubbletter"], m["tom_böjning_%"], m["böjning_eko"],
                f'{m["prio_1"]}/{m["prio_2"]}/{m["prio_3"]}', m["prio1_andel_andra_halvan"],
                dm.get("sekunder"), dm.get("kostnad"), r["dom"].get("kommentar", "")], 2):
                ws.cell(row=i, column=j, value=v)

        # --- fel som domaren hittat ---
        rad = len(rader) + 6
        ws.cell(row=rad, column=1, value="Fel domaren hittade").font = Font(bold=True, size=12)
        for j, k in enumerate(["modell", "glosa", "problem"], 1):
            c = ws.cell(row=rad + 1, column=j, value=k); c.font = RUBRIK; c.fill = FYLL
        rad += 2
        for r in rader:
            for fel in r["dom"].get("fel", []):
                ws.cell(row=rad, column=1, value=r["modell"])
                ws.cell(row=rad, column=2, value=fel.get("glosa"))
                ws.cell(row=rad, column=3, value=fel.get("problem"))
                rad += 1

        # --- listorna sida vid sida, två kolumner per modell ---
        start = rad + 2
        ws.cell(row=start, column=1, value="Glosorna").font = Font(bold=True, size=12)
        for i, r in enumerate(rader):
            c0 = 1 + i * 3
            c = ws.cell(row=start + 1, column=c0, value=r["modell"]); c.font = RUBRIK; c.fill = FYLL
            ws.cell(row=start + 1, column=c0 + 1, value="svenska").font = RUBRIK
            ws.cell(row=start + 1, column=c0 + 1).fill = FYLL
            for j, g in enumerate(r["data"].get("traffar", []), start + 2):
                ws.cell(row=j, column=c0, value=f'{g.get("prio","")} {g.get("ord","")}')
                ws.cell(row=j, column=c0 + 1, value=g.get("oversattning", ""))
        kolbredd(ws, [26, 8, 8, 9, 7, 9, 10, 9, 11, 13, 11, 11, 11, 10, 10, 60])
        ws.freeze_panes = "B4"

        for r in rader:
            sammanst.setdefault(r["modell"], []).append(
                {"kid": kid, "poang": r["poang"], "matt": r["matt"], "facit": r["facit"]})

    # --- sammanfattning ---
    ws = wb.create_sheet("Sammanfattning", 0)
    ws["A1"] = "Flippa – vilken modell ska generera lektioner?"; ws["A1"].font = Font(bold=True, size=14)
    kol = ["modell", "snittpoäng", "sämsta körning", "snitt täckning", "dubbletter totalt",
           "snitt sekunder", "kostnad hela testet $", "kostnad per 50-lektion $"]
    for j, k in enumerate(kol, 1):
        c = ws.cell(row=3, column=j, value=k); c.font = RUBRIK; c.fill = FYLL
    rader = []
    for modell, k in sammanst.items():
        poang = [x["poang"] for x in k if x["poang"] is not None]
        kostnader = [x["facit"].get("kostnad") or 0 for x in k]
        femtio = [x["facit"].get("kostnad") or 0 for x in k if x["kid"].endswith("-50")]
        sek = [x["facit"].get("sekunder") or 0 for x in k]
        rader.append([
            modell,
            round(sum(poang) / len(poang), 2) if poang else None,
            min(poang) if poang else None,
            round(sum(x["matt"]["täckning"] for x in k) / len(k), 2),
            sum(x["matt"]["dubbletter"] for x in k),
            round(sum(sek) / len(sek), 1) if sek else None,
            round(sum(kostnader), 4),
            round(sum(femtio) / len(femtio), 4) if femtio else None,
        ])
    rader.sort(key=lambda r: -(r[1] or 0))
    for i, r in enumerate(rader, 4):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=1).font = FET
    ws.cell(row=len(rader) + 6, column=1, value="Rekommendation").font = Font(bold=True, size=12)
    ws.cell(row=len(rader) + 7, column=1, value="(fylls i av domaren när alla körningar är dömda)")
    kolbredd(ws, [30, 12, 15, 15, 17, 14, 22, 24])
    ws.freeze_panes = "B4"

    wb.save(utfil)
    print("skrev", utfil, f"({len(wb.sheetnames)} flikar)")

if __name__ == "__main__":
    main()
